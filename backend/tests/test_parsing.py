import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from app.config import Settings
from app.errors import AppError
from app.parsing import extract_document

SAMPLES = Path(__file__).resolve().parents[2] / "samples"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.mark.parametrize("extension,mime", [("pdf", "application/pdf"), ("docx", DOCX)])
def test_valid_resume(extension, mime):
    text, _ = extract_document((SAMPLES / f"alex-morgan.{extension}").read_bytes(), f"cv.{extension}", mime, Settings())
    assert "Alex Morgan" in text
    assert "25%" in text
    assert "Volunteering" in text


@pytest.mark.parametrize("data,name,mime,code", [
    (b"", "cv.pdf", "application/pdf", "EMPTY_FILE"),
    (b"plain text", "cv.txt", "text/plain", "UNSUPPORTED_FILE"),
    (b"plain text", "cv.pdf", "application/pdf", "FILE_TYPE_MISMATCH"),
    (b"%PDF-corrupted", "cv.pdf", "application/pdf", "CORRUPT_DOCUMENT"),
    (b"pdf", "cv.docx", DOCX, "FILE_TYPE_MISMATCH"),
    (b"pdf", "cv.pdf", "image/png", "UNSUPPORTED_FILE"),
    (b"x" * (5 * 1024 * 1024 + 1), "cv.pdf", "application/pdf", "FILE_TOO_LARGE"),
], ids=["empty", "unsupported", "renamed-pdf", "corrupt-pdf", "renamed-docx", "bad-mime", "oversize"])
def test_bad_documents(data, name, mime, code):
    with pytest.raises(AppError) as exc:
        extract_document(data, name, mime, Settings())
    assert exc.value.code == code


def test_zip_renamed_as_docx():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("test.txt", "Not a Word document")
    with pytest.raises(AppError, match="not a DOCX"):
        extract_document(stream.getvalue(), "cv.docx", DOCX, Settings())


def test_blank_and_encrypted_pdf():
    for encrypted, code in [(False, "INSUFFICIENT_TEXT"), (True, "ENCRYPTED_PDF")]:
        stream = io.BytesIO()
        pdf = canvas.Canvas(stream, encrypt="secret" if encrypted else None)
        pdf.showPage(); pdf.save()
        with pytest.raises(AppError) as exc:
            extract_document(stream.getvalue(), "cv.pdf", "application/pdf", Settings())
        assert exc.value.code == code


def test_docx_table_reading_order():
    doc = Document()
    doc.add_paragraph("First paragraph with enough useful resume content for extraction.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table employer"
    table.cell(0, 1).text = "Table qualification"
    doc.add_paragraph("Last paragraph follows the table.")
    stream = io.BytesIO(); doc.save(stream)
    text, _ = extract_document(stream.getvalue(), "cv.docx", DOCX, Settings())
    assert text.index("First paragraph") < text.index("Table employer") < text.index("Last paragraph")


def test_resume_limit_never_truncates():
    with pytest.raises(AppError) as exc:
        extract_document((SAMPLES / "alex-morgan.pdf").read_bytes(), "cv.pdf", "application/pdf", Settings(max_resume_chars=100))
    assert exc.value.code == "RESUME_TOO_LONG"
