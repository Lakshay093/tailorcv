import io
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.errors import AppError
from app.schemas import Resume

ASSETS = Path(__file__).parent / "assets"


def blocks(resume: Resume) -> list[tuple[str, str]]:
    result = [("name", resume.contact.name)]
    contact = [resume.contact.email, resume.contact.phone, resume.contact.location, *resume.contact.links]
    result.append(("contact", " | ".join(x for x in contact if x)))
    if resume.summary:
        result.extend([("heading", "Summary"), ("body", resume.summary)])
    if resume.skills:
        result.extend([("heading", "Skills"), ("body", " · ".join(resume.skills))])
    if resume.experience:
        result.append(("heading", "Experience"))
        for entry in resume.experience:
            result.append(("role", " | ".join(x for x in [entry.title, entry.employer, entry.location] if x)))
            result.append(("date", " - ".join(x for x in [entry.start_date, entry.end_date] if x)))
            result.extend(("bullet", b.text) for b in entry.bullets)
    if resume.projects:
        result.append(("heading", "Projects"))
        for entry in resume.projects:
            result.append(("role", entry.name))
            if entry.description:
                result.append(("body", entry.description))
            result.extend(("bullet", b.text) for b in entry.bullets)
    if resume.education:
        result.append(("heading", "Education"))
        for entry in resume.education:
            result.append(("role", " | ".join([entry.qualification, entry.institution])))
            result.append(("date", entry.dates))
            result.extend(("body", d) for d in entry.details)
    if resume.certifications:
        result.append(("heading", "Certifications"))
        result.extend(("bullet", c) for c in resume.certifications)
    for extra in resume.additional_sections:
        result.append(("heading", extra.title))
        result.extend(("body", item) for item in extra.items)
    return [(kind, text) for kind, text in result if text]


def resume_text(resume: Resume) -> str:
    return "\n\n".join(("• " if kind == "bullet" else "") + text for kind, text in blocks(resume)) + "\n"


def pdf_bytes(resume: Resume) -> bytes:
    if "NotoSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("NotoSans", str(ASSETS / "NotoSans-Regular.ttf")))
        pdfmetrics.registerFont(TTFont("NotoSans-Bold", str(ASSETS / "NotoSans-Bold.ttf")))
        pdfmetrics.registerFontFamily("NotoSans", normal="NotoSans", bold="NotoSans-Bold", italic="NotoSans", boldItalic="NotoSans-Bold")
    # Fail explicitly for unsupported glyphs instead of silently exporting boxes.
    font = pdfmetrics.getFont("NotoSans")
    for _, text in blocks(resume):
        if any(ord(c) not in font.face.charToGlyph for c in text if not c.isspace()):
            raise AppError("UNSUPPORTED_PDF_GLYPH", "This PDF font does not support one or more characters. Download DOCX or text, or use a supported script.", 422)
    stream = io.BytesIO()
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("CVBody", parent=styles["Normal"], fontName="NotoSans", fontSize=9.5, leading=14.5, textColor=colors.HexColor("#202632"), spaceAfter=5, splitLongWords=True)
    style_map = {
        "body": normal,
        "bullet": ParagraphStyle("CVBullet", parent=normal, leftIndent=11, firstLineIndent=-9, spaceAfter=4),
        "name": ParagraphStyle("CVName", parent=normal, fontName="NotoSans-Bold", fontSize=22, leading=29, alignment=TA_CENTER, spaceAfter=5, keepWithNext=True),
        "contact": ParagraphStyle("CVContact", parent=normal, fontSize=8, leading=12, alignment=TA_CENTER, spaceAfter=12),
        "heading": ParagraphStyle("CVHeading", parent=normal, fontName="NotoSans-Bold", fontSize=11, leading=17, spaceBefore=11, spaceAfter=6, textColor=colors.HexColor("#3730a3"), keepWithNext=True),
        "role": ParagraphStyle("CVRole", parent=normal, fontName="NotoSans-Bold", spaceBefore=5, keepWithNext=True),
        "date": ParagraphStyle("CVDate", parent=normal, fontSize=8.5, textColor=colors.HexColor("#5d6470"), keepWithNext=True),
    }
    story = []
    for kind, text in blocks(resume):
        safe = escape(text).replace("\n", "<br/>")
        story.append(Paragraph(("• " if kind == "bullet" else "") + safe, style_map[kind]))
    if not story:
        story.append(Spacer(1, 1))
    SimpleDocTemplate(stream, pagesize=(8.27 * inch, 11.69 * inch), rightMargin=0.7 * inch, leftMargin=0.7 * inch, topMargin=0.6 * inch, bottomMargin=0.6 * inch, title=f"{resume.contact.name} - Resume", author="").build(story)
    return stream.getvalue()


def docx_bytes(resume: Resume) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10)
    base.paragraph_format.space_after = Pt(5)
    for kind, text in blocks(resume):
        p = doc.add_paragraph(style="List Bullet" if kind == "bullet" else "Normal")
        run = p.add_run(text)
        if kind in {"name", "heading", "role"}:
            run.bold = True
            p.paragraph_format.keep_with_next = True
        if kind == "name":
            p.alignment = 1
            run.font.size = Pt(22)
        if kind == "contact":
            p.alignment = 1
            run.font.size = Pt(8)
        if kind == "heading":
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor.from_string("3730A3")
            p.paragraph_format.space_before = Pt(12)
        if kind == "date":
            run.font.size = Pt(9)
            p.paragraph_format.keep_with_next = True
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def export_resume(resume: Resume, fmt: str) -> tuple[bytes, str]:
    if len(resume.model_dump_json()) > 200_000:
        raise AppError("RESUME_TOO_LONG", "The edited resume is too large to export.", 413)
    if fmt == "pdf":
        return pdf_bytes(resume), "application/pdf"
    if fmt == "docx":
        return docx_bytes(resume), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return resume_text(resume).encode("utf-8"), "text/plain; charset=utf-8"
